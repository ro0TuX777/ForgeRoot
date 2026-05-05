import sys
import os
import docx
from docx.shared import Pt

def write_md_to_doc(doc, md_file):
    with open(md_file, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    in_code_block = False
    code_content = []
    
    for line in lines:
        raw_line = line.replace('\n', '')
        
        if raw_line.startswith('```'):
            if in_code_block:
                p = doc.add_paragraph()
                run = p.add_run('\n'.join(code_content))
                run.font.name = 'Courier New'
                run.font.size = Pt(9)
                in_code_block = False
                code_content = []
            else:
                in_code_block = True
            continue
            
        if in_code_block:
            code_content.append(raw_line)
            continue
            
        clean_line = raw_line.replace('**', '').replace('__', '')
        
        if clean_line.startswith('# '):
            doc.add_heading(clean_line[2:], level=1)
        elif clean_line.startswith('## '):
            doc.add_heading(clean_line[3:], level=2)
        elif clean_line.startswith('### '):
            doc.add_heading(clean_line[4:], level=3)
        elif clean_line.startswith('---'):
            doc.add_paragraph('_' * 40)
        elif clean_line.strip() == '':
            pass
        else:
            doc.add_paragraph(clean_line)

def build_docx(md_path):
    if not os.path.exists(md_path):
        print(f"File not found: {md_path}")
        return
    
    doc = docx.Document()
    write_md_to_doc(doc, md_path)
    
    # Save as docx with the same name
    dir_name = os.path.dirname(md_path)
    base_name = os.path.basename(md_path).replace('.md', '.docx')
    docx_path = os.path.join(dir_name, base_name)
    
    doc.save(docx_path)
    print(f"Successfully wrote {docx_path}")

if __name__ == "__main__":
    if len(sys.argv) > 1:
        for md_file in sys.argv[1:]:
            build_docx(md_file)
    else:
        print("Usage: python md_to_docx.py <file1.md> <file2.md> ...")
